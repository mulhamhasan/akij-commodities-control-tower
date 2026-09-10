#!/usr/bin/env python3
"""GitHub Actions sync: read the register via Google Drive API and rebuild ctr-data.js.

Credentials are injected via environment variables (GitHub secrets).
"""
import io, json, os, re, urllib.request, urllib.parse

REGISTER_ID = os.environ.get("REGISTER_ID", "1me7vjHVAMV6Wazw4GnIzoPvvb3B_Rhz1CyDUY9FGLh4")

COMMODITY_ORDER = [
    'Wheat', 'Lentil', 'SBM / Soybean Meal', 'Corn', 'Maize', 'Soyabean',
    'Mustard', 'Yellow Peas', 'Chickpeas', 'Canola', 'Coal', 'General / Macro',
    'Palm Oil', 'Jute', 'Lysine/Threonine', 'Rice Bran', 'Groundnut', 'Turmeric'
]

RULES = [
    ('lysine', 'Lysine/Threonine'), ('threonine', 'Lysine/Threonine'),
    ('methionine', 'Lysine/Threonine'), ('choline', 'Lysine/Threonine'),
    ('monocalcium', 'Lysine/Threonine'), ('limestone', 'Lysine/Threonine'),
    ('sodium bicarbonate', 'Lysine/Threonine'), ('dcp', 'Lysine/Threonine'),
    ('rice br', 'Rice Bran'), ('dorb', 'Rice Bran'),
    ('turmeric', 'Turmeric'), ('groundnut', 'Groundnut'), ('jute', 'Jute'),
    ('palm', 'Palm Oil'), ('coal', 'Coal'),
    ('canola', 'Canola'), ('rapeseed', 'Canola'),
    ('chickpea', 'Chickpeas'),
    ('yellowpea', 'Yellow Peas'), ('yellow pea', 'Yellow Peas'),
    ('peas', 'Yellow Peas'), ('pulse', 'Yellow Peas'),
    ('mustard', 'Mustard'), ('lentil', 'Lentil'),
    ('soya bean meal', 'SBM / Soybean Meal'), ('soya meal', 'SBM / Soybean Meal'),
    ('soybean meal', 'SBM / Soybean Meal'), ('soy meal', 'SBM / Soybean Meal'),
    ('sbm', 'SBM / Soybean Meal'),
    ('soybean extraction', 'Soyabean'), ('soyabean extraction', 'Soyabean'),
    ('soybean', 'Soyabean'), ('soyabean', 'Soyabean'),
    ('corn', 'Corn'), ('wheat', 'Wheat'), ('cwrs', 'Wheat'), ('maize', 'Maize'),
]


def classify(item):
    s = " " + item.lower() + " "
    for k, com in RULES:
        if k in s:
            return com
    return "General / Macro"


def get_access_token():
    data = urllib.parse.urlencode({
        "client_id": os.environ["GOOGLE_CLIENT_ID"],
        "client_secret": os.environ["GOOGLE_CLIENT_SECRET"],
        "refresh_token": os.environ["GOOGLE_REFRESH_TOKEN"],
        "grant_type": "refresh_token",
    }).encode()
    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=data, method="POST")
    try:
        return json.load(urllib.request.urlopen(req, timeout=30))["access_token"]
    except urllib.error.HTTPError as e:
        print("token refresh failed HTTP %s: %s" % (e.code, e.read().decode()[:500]))
        raise


def download_xlsx(token):
    url = ("https://www.googleapis.com/drive/v3/files/%s/export"
           "?mimeType=application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" % REGISTER_ID)
    req = urllib.request.Request(url, headers={"Authorization": "Bearer " + token})
    return urllib.request.urlopen(req, timeout=60).read()


def href(cell):
    try:
        return cell.hyperlink.target if cell.hyperlink else None
    except Exception:
        return None


def extract_sheet(ws, category):
    rows = []
    hdr = None
    for r in range(1, 12):
        a = str(ws.cell(row=r, column=1).value or "")
        c = str(ws.cell(row=r, column=3).value or "")
        if "Report Type" in a or "Sub" in c:
            hdr = r
            break
    if hdr is None:
        return rows
    cur_group = None
    for r in range(hdr + 1, ws.max_row + 1):
        g = ws.cell(row=r, column=1).value
        if g is not None and str(g).strip():
            cur_group = str(g).strip()
        item = ws.cell(row=r, column=3).value
        if item is None or not str(item).strip():
            continue
        timeline = ws.cell(row=r, column=4).value
        link = None
        for c in range(5, ws.max_column + 1):
            h = href(ws.cell(row=r, column=c))
            if h:
                link = h
                break
        rows.append({
            "category": category,
            "group": cur_group,
            "item": str(item).strip(),
            "timeline": (str(timeline).strip() if timeline else ""),
            "link": link,
        })
    return rows


def detect_meetings(wb):
    nums = set()
    for ws in wb.worksheets:
        title = ws.title.lower()
        if not (("supply" in title) or ("demand" in title) or ("finance" in title)):
            continue
        for r in range(1, 11):
            for c in range(1, min(ws.max_column, 40) + 1):
                v = ws.cell(row=r, column=c).value
                if v and isinstance(v, str) and re.search(r"meeting", v, re.I):
                    m = re.search(r"(\d+)", v)
                    if m:
                        nums.add(int(m.group(1)))
    lst = sorted(nums)
    out = []
    for i, n in enumerate(lst):
        ord_ = "st" if n == 1 else "nd" if n == 2 else "rd" if n == 3 else "th"
        out.append({"n": n, "label": "%d%s Meeting" % (n, ord_), "status": "latest" if i == len(lst) - 1 else "past"})
    return out


def main():
    import openpyxl

    token = get_access_token()
    raw = download_xlsx(token)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)

    def find(frag):
        for s in wb.worksheets:
            if frag in s.title.lower():
                return s
        return None

    all_rows = []
    for cat, frag in (("supply", "supply"), ("demand", "demand"), ("finance", "finance")):
        ws = find(frag)
        if ws is not None:
            all_rows += extract_sheet(ws, cat)

    data = {}
    for row in all_rows:
        com = classify(row["item"])
        data.setdefault(com, {"supply": [], "demand": [], "finance": []})
        data[com][row["category"]].append({
            "item": row["item"],
            "timeline": row["timeline"],
            "link": row["link"],
            "group": row["group"],
        })

    ordered = {}
    for k in COMMODITY_ORDER:
        if k in data:
            ordered[k] = data[k]
    for k in data:
        if k not in ordered:
            ordered[k] = data[k]

    meetings = detect_meetings(wb)

    js = ("// Auto-generated from \"Source File\" register (Ch1 Supply / Ch2 Demand / Ch3 Finance).\n"
          "window.MEETINGS = " + json.dumps(meetings, ensure_ascii=False, separators=(",", ":")) + ";\n"
          "window.CTR_DATA = " + json.dumps(ordered, ensure_ascii=False, separators=(",", ":")) + ";\n")

    with open("ctr-data.js", "w") as f:
        f.write(js)
    print("generated ctr-data.js: %d items, %d commodities, %d meetings" % (len(all_rows), len(ordered), len(meetings)))


if __name__ == "__main__":
    main()
